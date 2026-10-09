import asyncio
from concurrent.futures import ThreadPoolExecutor
import io
import logging
import os
import re
import statistics
import sys
import tempfile
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from typing import Optional

from bson import ObjectId
from bson.errors import InvalidId
from dotenv import load_dotenv
from fastapi import FastAPI, Depends, HTTPException, Request, UploadFile, File, Form
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from jose import JWTError, jwt
from motor.motor_asyncio import AsyncIOMotorClient, AsyncIOMotorGridFSBucket
from passlib.context import CryptContext
from pydantic import BaseModel, EmailStr, Field, field_validator

try:
    import fitz
except Exception:  # pragma: no cover
    fitz = None

try:
    import cv2
    import numpy as np
except Exception:  # pragma: no cover
    cv2 = None
    np = None

try:
    from docx import Document as DocxDocument
except Exception:  # pragma: no cover
    DocxDocument = None

try:
    from paddleocr import PaddleOCR
except Exception:  # pragma: no cover
    PaddleOCR = None

load_dotenv()

MONGO_URI = os.getenv("MONGO_URI")
DB_NAME = os.getenv("DB_NAME", "Analysis-AI")
SECRET_KEY = os.getenv("SECRET_KEY", "dev-secret")
ALGORITHM = "HS256"
TOKEN_MINUTES = int(os.getenv("ACCESS_TOKEN_EXPIRE_MINUTES", "10080"))
COOKIE_NAME = "aia_token"
MAX_UPLOAD_BYTES = 25 * 1024 * 1024  # 25 MB
OCR_TIMEOUT_SECONDS = 90
OCR_LOG_TEXT = os.getenv("OCR_LOG_TEXT", "").strip().lower() in {
    "1",
    "true",
    "yes",
}

pwd_ctx = CryptContext(schemes=["bcrypt"], deprecated="auto")

app = FastAPI(title="AI Analysis")
app.mount("/static", StaticFiles(directory="static"), name="static")
templates = Jinja2Templates(directory="templates")

client: Optional[AsyncIOMotorClient] = None
db = None
bucket: Optional[AsyncIOMotorGridFSBucket] = None
ocr_engine = None
ocr_executor: Optional[ThreadPoolExecutor] = None
extraction_tasks: set[asyncio.Task] = set()
extraction_semaphore = asyncio.Semaphore(1)

# Configure logging to print to terminal at every step
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
    handlers=[
        logging.StreamHandler(sys.stdout),
    ],
    force=True,
)
logger = logging.getLogger(__name__)
print("[INIT] Logging configured - printing to terminal enabled")


@app.on_event("startup")
async def startup():
    global client, db, bucket
    print("[STARTUP] Initializing MongoDB connection...")
    logger.info("STARTUP: Initializing MongoDB connection to DB=%s", DB_NAME)
    client = AsyncIOMotorClient(MONGO_URI, serverSelectionTimeoutMS=20000)
    db = client[DB_NAME]
    bucket = AsyncIOMotorGridFSBucket(db, bucket_name="user_files")
    await db.users.create_index("email", unique=True)
    await db.documents.create_index([("user_id", 1), ("uploaded_at", -1)])
    await db.documents.create_index("extracted_text")
    print("[STARTUP] MongoDB connected and indexes created")
    logger.info("STARTUP: MongoDB connected and indexes created")
    async for doc in db.documents.find({"extraction_status": "pending"}):
        task = asyncio.create_task(
            _process_document_extraction(doc["_id"], doc["user_id"])
        )
        _schedule_extraction(task)


@app.on_event("shutdown")
async def shutdown():
    global ocr_executor
    if client:
        client.close()
    if ocr_executor is not None:
        ocr_executor.shutdown(wait=False)
        ocr_executor = None


# ---------------- models ----------------
class SignupIn(BaseModel):
    name: str = Field(min_length=2, max_length=80)
    email: EmailStr
    phone: str = Field(min_length=6, max_length=20)
    password: str = Field(min_length=6, max_length=128)
    confirm_password: str

    @field_validator("confirm_password")
    @classmethod
    def match(cls, v, info):
        if v != info.data.get("password"):
            raise ValueError("Passwords do not match")
        return v


class LoginIn(BaseModel):
    email: EmailStr
    password: str


class ProfileIn(BaseModel):
    name: Optional[str] = Field(default=None, max_length=80)
    phone: Optional[str] = Field(default=None, max_length=20)
    current_password: Optional[str] = None
    new_password: Optional[str] = Field(default=None, max_length=128)
    confirm_password: Optional[str] = None


class DocumentTextOut(BaseModel):
    id: str
    title: Optional[str]
    filename: Optional[str]
    content_type: Optional[str]
    size: int
    uploaded_at: Optional[str]
    extracted_text: Optional[str] = None


# ---------------- auth helpers ----------------
def make_token(user_id: str) -> str:
    payload = {"sub": user_id, "exp": datetime.now(timezone.utc) + timedelta(minutes=TOKEN_MINUTES)}
    return jwt.encode(payload, SECRET_KEY, algorithm=ALGORITHM)


def set_cookie(resp, token: str):
    resp.set_cookie(
        COOKIE_NAME, token, httponly=True, samesite="lax",
        max_age=TOKEN_MINUTES * 60, path="/",
    )


async def get_user_optional(request: Request):
    token = request.cookies.get(COOKIE_NAME)
    if not token:
        return None
    try:
        uid = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM]).get("sub")
        user = await db.users.find_one({"_id": ObjectId(uid)})
    except (JWTError, InvalidId, TypeError):
        return None
    return user


async def current_user(request: Request):
    user = await get_user_optional(request)
    if not user:
        raise HTTPException(status_code=401, detail="Not authenticated")
    return user


def public(user) -> dict:
    created = user.get("created_at")
    return {
        "id": str(user["_id"]),
        "name": user.get("name", ""),
        "email": user.get("email", ""),
        "phone": user.get("phone", ""),
        "created_at": created.isoformat() if created else None,
    }


def _get_file_ext(filename: Optional[str]) -> str:
    if not filename:
        return ""
    return os.path.splitext(filename)[1].lower()


def _get_ocr_engine():
    global ocr_engine
    print("[STEP 1] Calling _get_ocr_engine()")
    logger.info("STEP 1: Calling _get_ocr_engine()")
    if ocr_engine is not None:
        print("[STEP 1] OCR engine already cached, returning it")
        logger.info("STEP 1: OCR engine already cached")
        return ocr_engine
    if PaddleOCR is None:
        print("[ERROR] PaddleOCR is not installed")
        logger.error("ERROR: PaddleOCR is not installed")
        raise RuntimeError(
            "PaddleOCR is not installed in the active Python environment. "
            "Install the project requirements and restart the app."
        )
    print("[STEP 1] PaddleOCR is available, initializing engine...")
    logger.info("STEP 1: PaddleOCR is available, initializing engine")
    try:
        print("[STEP 1a] Attempting PaddleOCR initialization (variant 1)")
        logger.info("STEP 1a: Attempting PaddleOCR initialization (variant 1)")
        ocr_engine = PaddleOCR(
            lang="en",
            text_detection_model_name="PP-OCRv5_mobile_det",
            text_recognition_model_name="PP-OCRv5_mobile_rec",
            use_doc_orientation_classify=False,
            use_doc_unwarping=False,
            use_textline_orientation=False,
            text_det_limit_side_len=960,
            text_det_limit_type="max",
            device="cpu",
            enable_mkldnn=False,
        )
        print("[STEP 1a] PaddleOCR initialized (variant 1)")
        logger.info("STEP 1a: PaddleOCR initialized (variant 1)")
    except Exception as e:
        print(f"[ERROR] Failed to initialize PaddleOCR: {e}")
        logger.error("ERROR: Failed to initialize PaddleOCR: %s", e)
        raise
    print("[STEP 1] OCR engine initialization complete")
    logger.info("STEP 1: OCR engine initialization complete")
    return ocr_engine


@dataclass(frozen=True)
class _OCRDetection:
    text: str
    score: Optional[float] = None
    box: Optional[tuple[float, float, float, float]] = None

    @property
    def x(self) -> Optional[float]:
        return self.box[0] if self.box is not None else None

    @property
    def y(self) -> Optional[float]:
        return self.box[1] if self.box is not None else None

    @property
    def width(self) -> Optional[float]:
        return self.box[2] - self.box[0] if self.box is not None else None

    @property
    def height(self) -> Optional[float]:
        return self.box[3] - self.box[1] if self.box is not None else None


def _to_plain_list(value):
    if hasattr(value, "tolist"):
        value = value.tolist()
    if isinstance(value, (list, tuple)):
        return list(value)
    return []


def _box_bounds(value) -> Optional[tuple[float, float, float, float]]:
    points = _to_plain_list(value)
    if not points:
        return None
    if len(points) >= 4 and all(
        isinstance(coord, (int, float)) for coord in points[:4]
    ):
        x1, y1, x2, y2 = (float(coord) for coord in points[:4])
        return min(x1, x2), min(y1, y2), max(x1, x2), max(y1, y2)

    flattened = []
    for point in points:
        coordinates = _to_plain_list(point)
        if len(coordinates) >= 2 and all(
            isinstance(coord, (int, float)) for coord in coordinates[:2]
        ):
            flattened.append((float(coordinates[0]), float(coordinates[1])))
    if not flattened:
        return None
    xs, ys = zip(*flattened)
    return min(xs), min(ys), max(xs), max(ys)


def _collect_ocr_detections(result) -> list[_OCRDetection]:
    detections = []

    def collect(item):
        if isinstance(item, dict):
            texts = item.get("rec_texts")
            if texts:
                texts = _to_plain_list(texts)
                item_scores = _to_plain_list(item.get("rec_scores"))
                raw_boxes = item.get("rec_boxes")
                if raw_boxes is None:
                    raw_boxes = item.get("dt_polys")
                item_boxes = _to_plain_list(raw_boxes)
                for index, text in enumerate(texts):
                    if text is None:
                        continue
                    score = None
                    try:
                        if index < len(item_scores):
                            score = float(item_scores[index])
                    except (TypeError, ValueError):
                        pass
                    box = (
                        _box_bounds(item_boxes[index])
                        if index < len(item_boxes)
                        else None
                    )
                    detections.append(_OCRDetection(str(text), score, box))
            else:
                for value in item.values():
                    collect(value)
        elif isinstance(item, Iterable) and not isinstance(item, (str, bytes, bytearray)):
            if (
                isinstance(item, (list, tuple))
                and len(item) >= 2
                and isinstance(item[1], (list, tuple))
                and item[1]
                and isinstance(item[1][0], str)
            ):
                score = None
                if len(item[1]) > 1:
                    try:
                        score = float(item[1][1])
                    except (TypeError, ValueError):
                        pass
                detections.append(
                    _OCRDetection(
                        item[1][0],
                        score,
                        _box_bounds(item[0]),
                    )
                )
            else:
                for value in item:
                    collect(value)

    collect(result)
    return detections


def _is_watermark_detection(
    detection: _OCRDetection,
    detections: list[_OCRDetection],
    image_size: tuple[int, int],
) -> bool:
    if detection.box is None:
        return False
    width, height = image_size
    x1, y1, _, y2 = detection.box
    text = detection.text.strip()
    box_height = y2 - y1
    if (
        x1 < width * 0.58
        or y1 < height * 0.68
        or box_height > height * 0.04
        or len(text) > 32
    ):
        return False

    camera_cue = re.compile(
        r"\b(?:camera|dual\s+camera|shot\s+on|captured\s+on|watermark)\b",
        re.IGNORECASE,
    )
    if camera_cue.search(text):
        return True

    # Device/model text is commonly split into a second nearby OCR box.
    has_nearby_camera_cue = any(
        other is not detection
        and other.box is not None
        and camera_cue.search(other.text)
        and other.box[0] >= width * 0.58
        and other.box[1] >= height * 0.68
        and abs((other.box[1] + other.box[3]) / 2 - (y1 + y2) / 2)
        <= height * 0.09
        for other in detections
    )
    return (
        has_nearby_camera_cue
        and text.isascii()
        and text.upper() == text
        and any(char.isalpha() for char in text)
        and len(text.split()) <= 6
    )


def _coordinate_order_rows(
    detections: list[_OCRDetection],
) -> list[list[_OCRDetection]]:
    positioned = [detection for detection in detections if detection.box is not None]
    unpositioned = [detection for detection in detections if detection.box is None]
    if not positioned:
        return [[detection] for detection in unpositioned]

    heights = [
        max(1.0, detection.box[3] - detection.box[1])
        for detection in positioned
    ]
    typical_height = statistics.median(heights)
    rows = []
    for detection in sorted(
        positioned,
        key=lambda item: (item.box[1], item.box[0]),
    ):
        box = detection.box
        center_y = (box[1] + box[3]) / 2
        detection_height = max(1.0, box[3] - box[1])
        candidates = []
        for row in rows:
            center_distance = abs(center_y - row["center_y"])
            row_height = row["height"]
            tolerance = (
                0.5 * min(detection_height, row_height)
                + 0.25 * typical_height
            )
            if center_distance <= tolerance:
                candidates.append(
                    (
                        center_distance / max(1.0, typical_height),
                        row,
                        center_y,
                        detection_height,
                    )
                )

        if not candidates:
            rows.append(
                {
                    "center_y": center_y,
                    "height": detection_height,
                    "detections": [detection],
                }
            )
            continue

        _, row, _, _ = min(candidates, key=lambda candidate: candidate[0])
        row["detections"].append(detection)
        centers = [
            (item.box[1] + item.box[3]) / 2 for item in row["detections"]
        ]
        row["center_y"] = statistics.median(centers)
        row["height"] = statistics.median(
            max(1.0, item.box[3] - item.box[1])
            for item in row["detections"]
        )

    ordered_rows = [
        sorted(row["detections"], key=lambda item: item.box[0])
        for row in sorted(rows, key=lambda item: item["center_y"])
    ]
    if unpositioned:
        logger.warning(
            "OCR returned %d detection(s) without coordinates; preserving them "
            "after coordinate-ordered detections",
            len(unpositioned),
        )
        ordered_rows.extend([[detection] for detection in unpositioned])
    return ordered_rows


def _format_ocr_detections(
    detections: list[_OCRDetection],
    image_size: Optional[tuple[int, int]] = None,
) -> tuple[str, list[float]]:
    raw_text = "\n".join(detection.text for detection in detections)
    if OCR_LOG_TEXT:
        logger.info("OCR raw output:\n%s", raw_text)
    else:
        logger.debug("OCR raw output:\n%s", raw_text)
    raw_ordered_rows = _coordinate_order_rows(detections)
    coordinate_ordered_text = "\n".join(
        " ".join(detection.text for detection in row)
        for row in raw_ordered_rows
    )
    if OCR_LOG_TEXT:
        logger.info("OCR coordinate-ordered output:\n%s", coordinate_ordered_text)
    else:
        logger.debug(
            "OCR coordinate-ordered output:\n%s",
            coordinate_ordered_text,
        )

    retained = []
    for detection in detections:
        text = "".join(
            character for character in detection.text if character.isprintable()
        )
        text = " ".join(text.split())
        if not text or not any(character.isalnum() for character in text):
            continue
        if detection.score is not None and detection.score < 0.35:
            text = "[unclear]"
        elif (
            len(text) == 1
            and not text.isascii()
            and detection.score is not None
            and detection.score < 0.55
        ):
            text = "[unclear]"
        if image_size and _is_watermark_detection(detection, detections, image_size):
            continue
        retained.append(_OCRDetection(text, detection.score, detection.box))

    # Remove only duplicate detections occupying the same location. Identical
    # values in different parts of a document remain untouched.
    unique = []
    for detection in retained:
        duplicate = False
        if detection.box is not None:
            for existing in unique:
                if existing.text != detection.text or existing.box is None:
                    continue
                ax1, ay1, ax2, ay2 = detection.box
                bx1, by1, bx2, by2 = existing.box
                intersection = max(0.0, min(ax2, bx2) - max(ax1, bx1)) * max(
                    0.0, min(ay2, by2) - max(ay1, by1)
                )
                area_a = max(0.0, ax2 - ax1) * max(0.0, ay2 - ay1)
                area_b = max(0.0, bx2 - bx1) * max(0.0, by2 - by1)
                if intersection / max(1.0, min(area_a, area_b)) >= 0.8:
                    duplicate = True
                    break
        if not duplicate:
            unique.append(detection)

    ordered_rows = _coordinate_order_rows(unique)

    cleaned_text = "\n".join(
        " ".join(detection.text for detection in row) for row in ordered_rows
    )
    if OCR_LOG_TEXT:
        logger.info("OCR cleaned output:\n%s", cleaned_text)
    else:
        logger.debug("OCR cleaned output:\n%s", cleaned_text)
    scores = [
        detection.score
        for row in ordered_rows
        for detection in row
        if detection.score is not None
    ]
    logger.info(
        "OCR postprocessing: detections=%d retained=%d output_chars=%d",
        len(detections),
        sum(len(row) for row in ordered_rows),
        len(cleaned_text),
    )
    return cleaned_text, scores


def _read_ocr_result(
    result,
    image_size: Optional[tuple[int, int]] = None,
) -> tuple[str, list[float]]:
    detections = _collect_ocr_detections(result)
    return _format_ocr_detections(detections, image_size)


def _read_ocr_text(result) -> str:
    return _read_ocr_result(result)[0]


def _preprocess_image_variants(data: bytes) -> list[bytes]:
    print("[STEP 2] Preprocessing image variants for OCR...")
    logger.info("STEP 2: Preprocessing image variants for OCR")
    if cv2 is None or np is None:
        print("[ERROR] OpenCV/numpy not installed")
        logger.error("ERROR: OpenCV/numpy not installed")
        raise RuntimeError(
            "OpenCV is not installed in the active Python environment. "
            "Install the project requirements and restart the app."
        )

    print("[STEP 2a] Decoding image from bytes")
    logger.info("STEP 2a: Decoding image from bytes")
    image = cv2.imdecode(np.frombuffer(data, dtype=np.uint8), cv2.IMREAD_COLOR)
    if image is None:
        print("[ERROR] Image could not be decoded")
        logger.error("ERROR: Image could not be decoded")
        raise ValueError("The uploaded image could not be decoded")
    print(f"[STEP 2a] Image decoded, shape={image.shape}")
    logger.info("STEP 2a: Image decoded, shape=%s", image.shape)

    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    code_regions = []
    try:
        qr_detected, qr_points = cv2.QRCodeDetector().detect(gray)
        if qr_detected and qr_points is not None:
            code_regions.append(qr_points)
    except Exception:
        logger.warning("QR-code detection failed; continuing without QR masking", exc_info=True)
    try:
        barcode_detector = cv2.barcode_BarcodeDetector()
        barcode_detected, barcode_points = barcode_detector.detect(image)
        if barcode_detected and barcode_points is not None:
            code_regions.extend(barcode_points)
    except Exception:
        logger.warning(
            "Barcode detection failed; continuing without barcode masking",
            exc_info=True,
        )

    for points in code_regions:
        polygon = np.asarray(points, dtype=np.int32).reshape(-1, 2)
        if len(polygon) >= 3:
            cv2.fillConvexPoly(image, cv2.convexHull(polygon), (255, 255, 255))
    if code_regions:
        logger.info("Masked %d QR/barcode region(s) before OCR", len(code_regions))

    height, width = image.shape[:2]
    scale = min(1.5, 1800 / max(height, width))
    if scale > 1:
        print(f"[STEP 2b] Resizing image with scale={scale:.2f}")
        logger.info("STEP 2b: Resizing image with scale=%.2f", scale)
        image = cv2.resize(
            image,
            (round(width * scale), round(height * scale)),
            interpolation=cv2.INTER_CUBIC,
        )
        print(f"[STEP 2b] Image resized to {image.shape}")
        logger.info("STEP 2b: Image resized to %s", image.shape)

    print("[STEP 2c] Converting to grayscale, denoising, enhancing")
    logger.info("STEP 2c: Converting to grayscale, denoising, enhancing")
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    denoised = cv2.medianBlur(gray, 3)
    enhanced = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8)).apply(denoised)
    enhanced = cv2.addWeighted(
        enhanced,
        1.15,
        cv2.GaussianBlur(enhanced, (0, 0), 1.0),
        -0.15,
        0,
    )
    print("[STEP 2c] Preprocessing completed")
    logger.info("STEP 2c: Preprocessing completed")

    variants = []
    variants_to_encode = (cv2.cvtColor(enhanced, cv2.COLOR_GRAY2BGR),)
    for idx, variant in enumerate(variants_to_encode):
        print(f"[STEP 2d] Encoding variant {idx+1}/{len(variants_to_encode)} to PNG")
        logger.info("STEP 2d: Encoding variant %d/%d to PNG", idx+1, len(variants_to_encode))
        success, encoded = cv2.imencode(".png", variant)
        if not success:
            print(f"[ERROR] Failed to encode variant {idx+1}")
            logger.error("ERROR: Failed to encode variant %d", idx+1)
            raise RuntimeError("OpenCV could not encode an OCR preprocessing variant")
        variants.append(encoded.tobytes())
        print(f"[STEP 2d] Variant {idx+1} encoded ({len(variants[-1])} bytes)")
        logger.info("STEP 2d: Variant %d encoded (%d bytes)", idx+1, len(variants[-1]))
    print(f"[STEP 2] Generated {len(variants)} image variants")
    logger.info("STEP 2: Generated %d image variants", len(variants))
    return variants


def _run_ocr(engine, data: bytes, suffix: str) -> tuple[str, list[float]]:
    print(f"[STEP 3] Running OCR on temporary file (suffix={suffix}, size={len(data)} bytes)")
    logger.info("STEP 3: Running OCR on temp file (suffix=%s, size=%d bytes)", suffix, len(data))
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        tmp.write(data)
        tmp_path = tmp.name
    print(f"[STEP 3] Temp file written: {tmp_path}")
    logger.info("STEP 3: Temp file written: %s", tmp_path)
    try:
        if hasattr(engine, "predict"):
            print("[STEP 3a] Using engine.predict()")
            logger.info("STEP 3a: Using engine.predict()")
            result = engine.predict(input=tmp_path)
        else:
            print("[STEP 3a] Using engine.ocr()")
            logger.info("STEP 3a: Using engine.ocr()")
            result = engine.ocr(tmp_path, cls=True)
        print("[STEP 3a] OCR call completed, parsing results...")
        logger.info("STEP 3a: OCR call completed, parsing results")
        decoded = cv2.imdecode(np.frombuffer(data, dtype=np.uint8), cv2.IMREAD_COLOR)
        image_size = (decoded.shape[1], decoded.shape[0]) if decoded is not None else None
        text, scores = _read_ocr_result(result, image_size)
        print(f"[STEP 3a] Parsed OCR result: {len(text)} chars, {len(scores)} scores")
        logger.info("STEP 3a: Parsed OCR result: %d chars, %d scores", len(text), len(scores))
        return text, scores
    except Exception as e:
        print(f"[ERROR] OCR execution failed: {e}")
        logger.error("ERROR: OCR execution failed: %s", e)
        raise
    finally:
        try:
            os.unlink(tmp_path)
            print(f"[STEP 3] Cleaned up temp file: {tmp_path}")
            logger.info("STEP 3: Cleaned up temp file: %s", tmp_path)
        except Exception as e:
            print(f"[WARNING] Failed to cleanup temp file: {e}")
            logger.warning("WARNING: Failed to cleanup temp file: %s", e)


def _ocr_image_bytes(data: bytes, suffix: str = ".png") -> str:
    print(f"[STEP 4] OCR image bytes called (size={len(data)} bytes, suffix={suffix})")
    logger.info("STEP 4: OCR image bytes called (size=%d bytes, suffix=%s)", len(data), suffix)
    engine = _get_ocr_engine()
    variants = _preprocess_image_variants(data)
    candidates = []
    errors = []
    for index, variant in enumerate(variants):
        print(f"[STEP 4a] Processing OCR variant {index+1}/{len(variants)}")
        logger.info("STEP 4a: Processing OCR variant %d/%d", index+1, len(variants))
        try:
            text, scores = _run_ocr(engine, variant, ".png")
        except Exception as exc:
            print(f"[WARNING] OCR failed for variant {index+1}: {exc}")
            logger.warning("OCR failed for preprocessing variant %d: %s", index, exc)
            errors.append(exc)
            continue
        if text:
            average_confidence = sum(scores) / len(scores) if scores else -1.0
            print(f"[STEP 4a] Variant {index+1} produced text: {len(text)} chars, avg_conf={average_confidence:.3f}")
            logger.info("STEP 4a: Variant %d produced text: %d chars, avg_conf=%.3f", index+1, len(text), average_confidence)
            candidates.append((average_confidence, len(text), text))
        else:
            print(f"[STEP 4a] Variant {index+1} produced NO text")
            logger.info("STEP 4a: Variant %d produced NO text", index+1)

    if not candidates and errors:
        print(f"[ERROR] OCR failed for all variants, last error: {errors[-1]}")
        logger.error("ERROR: OCR failed for all image variants: %s", errors[-1])
        raise RuntimeError(f"OCR failed for all image variants: {errors[-1]}") from errors[-1]
    if not candidates:
        print("[STEP 4] No candidates found, returning empty string")
        logger.info("STEP 4: No candidates found, returning empty string")
        return ""
    best = max(candidates, key=lambda candidate: (candidate[0], candidate[1]))
    print(f"[STEP 4] Selected best variant with conf={best[0]:.3f}, len={best[1]}, returning {len(best[2])} chars")
    logger.info("STEP 4: Selected best variant with conf=%.3f, len=%d, returning %d chars", best[0], best[1], len(best[2]))
    return best[2]


def _extract_text_from_pdf_bytes(data: bytes) -> str:
    print(f"[STEP 5] Extracting text from PDF bytes (size={len(data)} bytes)")
    logger.info("STEP 5: Extracting text from PDF bytes (size=%d bytes)", len(data))
    if fitz is None:
        print("[ERROR] PyMuPDF is not installed")
        logger.error("ERROR: PyMuPDF is not installed")
        raise RuntimeError("PyMuPDF is not installed")
    text_chunks = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        num_pages = len(doc)
        print(f"[STEP 5] PDF opened, {num_pages} pages")
        logger.info("STEP 5: PDF opened, %d pages", num_pages)
        for i, page in enumerate(doc):
            print(f"[STEP 5a] Processing page {i+1}/{num_pages}")
            logger.info("STEP 5a: Processing page %d/%d", i+1, num_pages)
            text = page.get_text("text")
            if text.strip():
                print(f"[STEP 5a] Page {i+1} has embedded text ({len(text)} chars)")
                logger.info("STEP 5a: Page %d has embedded text (%d chars)", i+1, len(text))
                text_chunks.append(text)
            else:
                print(f"[STEP 5a] Page {i+1} has no embedded text - rendering to image for OCR (200 DPI)")
                logger.info("STEP 5a: Page %d has no embedded text - rendering to image for OCR (200 DPI)", i+1)
                page_image = page.get_pixmap(dpi=200, alpha=False).tobytes("png")
                print(f"[STEP 5a] Page {i+1} rendered to image ({len(page_image)} bytes)")
                logger.info("STEP 5a: Page %d rendered to image (%d bytes)", i+1, len(page_image))
                ocr_text = _ocr_image_bytes(page_image)
                print(f"[STEP 5a] Page {i+1} OCR returned {len(ocr_text)} chars")
                logger.info("STEP 5a: Page %d OCR returned %d chars", i+1, len(ocr_text))
                text_chunks.append(ocr_text)
    result = "\n".join(text_chunks)
    print(f"[STEP 5] PDF extraction complete, total {len(result)} chars from {len(text_chunks)} chunks")
    logger.info("STEP 5: PDF extraction complete, total %d chars from %d chunks", len(result), len(text_chunks))
    return result


def _extract_text_from_docx_bytes(data: bytes) -> str:
    print(f"[STEP 6] Extracting text from DOCX bytes (size={len(data)} bytes)")
    logger.info("STEP 6: Extracting text from DOCX bytes (size=%d bytes)", len(data))
    if DocxDocument is None:
        print("[ERROR] python-docx is not installed")
        logger.error("ERROR: python-docx is not installed")
        raise RuntimeError("python-docx is not installed")
    doc = DocxDocument(io.BytesIO(data))
    paragraph_text = []
    para_count = len(doc.paragraphs)
    print(f"[STEP 6] DOCX opened with {para_count} paragraphs")
    logger.info("STEP 6: DOCX opened with %d paragraphs", para_count)
    for i, paragraph in enumerate(doc.paragraphs):
        if i < 5 or i % 50 == 0:  # log first few and periodically
            print(f"[STEP 6a] Paragraph {i+1}/{para_count}: {len(paragraph.text)} chars")
            logger.debug("STEP 6a: Paragraph %d/%d: %d chars", i+1, para_count, len(paragraph.text))
        paragraph_text.append(paragraph.text)
    result = "\n".join(paragraph_text)
    print(f"[STEP 6] DOCX extraction complete, total {len(result)} chars")
    logger.info("STEP 6: DOCX extraction complete, total %d chars", len(result))
    return result


def _extract_raw_text_for_upload(filename: Optional[str], file_bytes: bytes, content_type: Optional[str]) -> str:
    print(f"[STEP 7] _extract_raw_text_for_upload called - filename={filename}, size={len(file_bytes)} bytes, content_type={content_type}")
    logger.info("STEP 7: _extract_raw_text_for_upload called - filename=%s, size=%d bytes, content_type=%s", filename, len(file_bytes), content_type)
    ext = _get_file_ext(filename)
    lower_type = (content_type or "").lower()
    print(f"[STEP 7] Detected ext={ext}, lower_type={lower_type}")
    logger.info("STEP 7: Detected ext=%s, lower_type=%s", ext, lower_type)

    if ext == ".pdf" or lower_type == "application/pdf":
        print("[STEP 7] Format is PDF -> calling PDF extractor")
        logger.info("STEP 7: Format is PDF -> calling PDF extractor")
        result = _extract_text_from_pdf_bytes(file_bytes)
        print(f"[STEP 7] PDF extraction returned {len(result)} chars (raw text)")
        logger.info("STEP 7: PDF extraction returned %d chars (raw text)", len(result))
        return result
    if ext == ".docx" or lower_type in {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/docx",
    }:
        print("[STEP 7] Format is DOCX -> calling DOCX extractor")
        logger.info("STEP 7: Format is DOCX -> calling DOCX extractor")
        result = _extract_text_from_docx_bytes(file_bytes)
        print(f"[STEP 7] DOCX extraction returned {len(result)} chars (raw text)")
        logger.info("STEP 7: DOCX extraction returned %d chars (raw text)", len(result))
        return result
    if ext in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"} or lower_type.startswith("image/"):
        print("[STEP 7] Format is IMAGE -> calling OCR extractor")
        logger.info("STEP 7: Format is IMAGE -> calling OCR extractor")
        result = _ocr_image_bytes(file_bytes, ext or ".png")
        print(f"[STEP 7] IMAGE OCR returned {len(result)} chars (raw text)")
        logger.info("STEP 7: IMAGE OCR returned %d chars (raw text)", len(result))
        return result
    print("[STEP 7] Unsupported format - returning empty string")
    logger.info("STEP 7: Unsupported format - returning empty string")
    return ""


def _get_ocr_executor() -> ThreadPoolExecutor:
    global ocr_executor
    if ocr_executor is None:
        ocr_executor = ThreadPoolExecutor(
            max_workers=1,
            thread_name_prefix="document-ocr",
        )
    return ocr_executor


async def _extract_with_timeout(
    filename: Optional[str],
    file_bytes: bytes,
    content_type: Optional[str],
) -> str:
    try:
        loop = asyncio.get_running_loop()
        return await asyncio.wait_for(
            loop.run_in_executor(
                _get_ocr_executor(),
                _extract_raw_text_for_upload,
                filename,
                file_bytes,
                content_type,
            ),
            timeout=OCR_TIMEOUT_SECONDS,
        )
    except asyncio.TimeoutError as exc:
        raise TimeoutError(
            f"Text extraction exceeded {OCR_TIMEOUT_SECONDS} seconds"
        ) from exc


async def _process_document_extraction(doc_id: ObjectId, user_id: ObjectId):
    try:
        doc = await db.documents.find_one({"_id": doc_id, "user_id": user_id})
        if not doc:
            return
        stream = await bucket.open_download_stream(doc["gridfs_id"])
        raw = await stream.read()
        async with extraction_semaphore:
            await db.documents.update_one(
                {"_id": doc_id, "user_id": user_id},
                {"$set": {"extraction_started_at": datetime.now(timezone.utc)}},
            )
            extracted_text = await _extract_with_timeout(
                doc.get("filename"),
                raw,
                doc.get("content_type"),
            )
        status = "completed" if extracted_text else "no_text"
        error = None
    except Exception as exc:
        extracted_text = ""
        status = "failed"
        error = str(exc)
        logger.exception("Text extraction failed for document %s", doc_id)

    await db.documents.update_one(
        {"_id": doc_id, "user_id": user_id},
        {
            "$set": {
                "extracted_text": extracted_text,
                "extraction_status": status,
                "extraction_error": error,
                "extraction_completed_at": datetime.now(timezone.utc),
            }
        },
    )


async def _process_uploaded_document_extraction(
    doc_id: ObjectId,
    user_id: ObjectId,
    filename: Optional[str],
    content_type: Optional[str],
    raw: bytes,
):
    try:
        logger.info("Background extraction started for document %s", doc_id)
        async with extraction_semaphore:
            await db.documents.update_one(
                {"_id": doc_id, "user_id": user_id},
                {"$set": {"extraction_started_at": datetime.now(timezone.utc)}},
            )
            extracted_text = await _extract_with_timeout(
                filename, raw, content_type
            )
        status = "completed" if extracted_text else "no_text"
        error = None
        logger.info(
            "Background extraction finished for document %s: status=%s, chars=%d",
            doc_id,
            status,
            len(extracted_text),
        )
    except Exception as exc:
        extracted_text = ""
        status = "failed"
        error = str(exc)
        logger.exception("Background extraction failed for document %s", doc_id)

    await db.documents.update_one(
        {"_id": doc_id, "user_id": user_id},
        {
            "$set": {
                "extracted_text": extracted_text,
                "extraction_status": status,
                "extraction_error": error,
                "extraction_completed_at": datetime.now(timezone.utc),
            }
        },
    )


def _schedule_extraction(task: asyncio.Task):
    extraction_tasks.add(task)

    def on_done(completed: asyncio.Task):
        extraction_tasks.discard(completed)
        if completed.cancelled():
            logger.warning("Background document extraction task was cancelled")
            return
        error = completed.exception()
        if error:
            logger.error(
                "Background document extraction task crashed",
                exc_info=(type(error), error, error.__traceback__),
            )

    task.add_done_callback(on_done)


# ---------------- pages ----------------
@app.get("/")
async def root(request: Request):
    user = await get_user_optional(request)
    return RedirectResponse("/dashboard" if user else "/login")


@app.get("/login", response_class=HTMLResponse)
async def login_page(request: Request):
    if await get_user_optional(request):
        return RedirectResponse("/dashboard")
    return templates.TemplateResponse("auth.html", {"request": request})


@app.get("/dashboard", response_class=HTMLResponse)
async def dashboard_page(request: Request):
    user = await get_user_optional(request)
    if not user:
        return RedirectResponse("/login")
    return templates.TemplateResponse("dashboard.html", {"request": request, "user": public(user)})


# ---------------- auth api ----------------
@app.post("/api/signup")
async def signup(data: SignupIn):
    if await db.users.find_one({"email": data.email.lower()}):
        raise HTTPException(status_code=400, detail="Email already registered")
    doc = {
        "name": data.name.strip(),
        "email": data.email.lower(),
        "phone": data.phone.strip(),
        "password_hash": pwd_ctx.hash(data.password),
        "created_at": datetime.now(timezone.utc),
    }
    res = await db.users.insert_one(doc)
    resp = JSONResponse({"ok": True, "redirect": "/dashboard"})
    set_cookie(resp, make_token(str(res.inserted_id)))
    return resp


@app.post("/api/login")
async def login(data: LoginIn):
    user = await db.users.find_one({"email": data.email.lower()})
    if not user or not pwd_ctx.verify(data.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Invalid email or password")
    resp = JSONResponse({"ok": True, "redirect": "/dashboard"})
    set_cookie(resp, make_token(str(user["_id"])))
    return resp


@app.post("/api/logout")
async def logout():
    resp = JSONResponse({"ok": True, "redirect": "/login"})
    resp.delete_cookie(COOKIE_NAME, path="/")
    return resp


@app.get("/api/me")
async def me(user=Depends(current_user)):
    return public(user)


@app.put("/api/me")
async def update_me(data: ProfileIn, user=Depends(current_user)):
    updates = {}
    if data.name and data.name.strip():
        updates["name"] = data.name.strip()
    if data.phone and data.phone.strip():
        updates["phone"] = data.phone.strip()

    if data.new_password:
        if not data.current_password or not pwd_ctx.verify(data.current_password, user["password_hash"]):
            raise HTTPException(status_code=400, detail="Current password is incorrect")
        if len(data.new_password) < 6:
            raise HTTPException(status_code=400, detail="New password must be at least 6 characters")
        if data.new_password != data.confirm_password:
            raise HTTPException(status_code=400, detail="New passwords do not match")
        updates["password_hash"] = pwd_ctx.hash(data.new_password)

    if not updates:
        raise HTTPException(status_code=400, detail="Nothing to update")

    updates["updated_at"] = datetime.now(timezone.utc)
    await db.users.update_one({"_id": user["_id"]}, {"$set": updates})
    fresh = await db.users.find_one({"_id": user["_id"]})
    return {"ok": True, "user": public(fresh)}


# ---------------- stats ----------------
@app.get("/api/stats")
async def stats(user=Depends(current_user)):
    uid = user["_id"]
    total = await db.documents.count_documents({"user_id": uid})

    agg = await db.documents.aggregate([
        {"$match": {"user_id": uid}},
        {"$group": {"_id": None, "bytes": {"$sum": "$size"}}},
    ]).to_list(1)
    total_bytes = agg[0]["bytes"] if agg else 0

    week_ago = datetime.now(timezone.utc) - timedelta(days=7)
    recent = await db.documents.count_documents({"user_id": uid, "uploaded_at": {"$gte": week_ago}})

    types = await db.documents.aggregate([
        {"$match": {"user_id": uid}},
        {"$group": {"_id": "$content_type", "count": {"$sum": 1}}},
        {"$sort": {"count": -1}},
        {"$limit": 6},
    ]).to_list(6)

    daily = await db.documents.aggregate([
        {"$match": {"user_id": uid, "uploaded_at": {"$gte": week_ago}}},
        {"$group": {
            "_id": {"$dateToString": {"format": "%Y-%m-%d", "date": "$uploaded_at"}},
            "count": {"$sum": 1},
        }},
        {"$sort": {"_id": 1}},
    ]).to_list(10)

    return {
        "total_documents": total,
        "total_bytes": total_bytes,
        "recent_7_days": recent,
        "distinct_types": len(types),
        "by_type": [{"type": t["_id"] or "unknown", "count": t["count"]} for t in types],
        "daily": {d["_id"]: d["count"] for d in daily},
    }


# ---------------- documents ----------------
@app.post("/api/upload")
async def upload(
    file: UploadFile = File(...),
    title: str = Form(""),
    user=Depends(current_user),
):
    raw = await file.read()
    if not raw:
        raise HTTPException(status_code=400, detail="Empty file")
    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="File too large (max 25 MB)")

    uid = user["_id"]
    gid = await bucket.upload_from_stream(
        file.filename, raw,
        metadata={"user_id": str(uid), "content_type": file.content_type},
    )

    ext = _get_file_ext(file.filename)
    lower_type = (file.content_type or "").lower()
    supported_format = (
        ext in {".pdf", ".docx", ".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"}
        or lower_type == "application/pdf"
        or lower_type in {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/docx",
        }
        or lower_type.startswith("image/")
    )

    logger.info(
        "Upload received: filename=%s, size=%d, content_type=%s",
        file.filename,
        len(raw),
        file.content_type,
    )
    extraction_status = "pending" if supported_format else "unsupported"

    doc = {
        "user_id": uid,
        "title": (title or "").strip() or file.filename,
        "filename": file.filename,
        "content_type": file.content_type or "application/octet-stream",
        "size": len(raw),
        "gridfs_id": gid,
        "uploaded_at": datetime.now(timezone.utc),
        "extracted_text": "",
        "extraction_status": extraction_status,
        "extraction_error": None,
        "extraction_started_at": None,
    }
    res = await db.documents.insert_one(doc)
    if supported_format:
        _schedule_extraction(
            asyncio.create_task(
                _process_uploaded_document_extraction(
                    res.inserted_id,
                    uid,
                    file.filename,
                    file.content_type,
                    raw,
                )
            )
        )
    logger.info("Upload saved: document_id=%s, extraction_status=%s", res.inserted_id, extraction_status)
    return {
        "ok": True,
        "id": str(res.inserted_id),
        "title": doc["title"],
        "size": doc["size"],
        "extracted_text": "",
        "extraction_status": extraction_status,
        "extraction_error": None,
    }


@app.get("/api/documents")
async def list_documents(user=Depends(current_user), q: str = "", limit: int = 100):
    query = {"user_id": user["_id"]}
    if q.strip():
        query["title"] = {"$regex": q.strip(), "$options": "i"}
    cur = db.documents.find(query).sort("uploaded_at", -1).limit(min(limit, 200))
    out = []
    async for d in cur:
        out.append({
            "id": str(d["_id"]),
            "title": d.get("title"),
            "filename": d.get("filename"),
            "content_type": d.get("content_type"),
            "size": d.get("size", 0),
            "uploaded_at": d["uploaded_at"].isoformat() if d.get("uploaded_at") else None,
            "has_extracted_text": bool(d.get("extracted_text")),
            "extraction_status": d.get("extraction_status", "pending"),
        })
    return {"documents": out}


async def _owned_doc(doc_id: str, user):
    try:
        oid = ObjectId(doc_id)
    except InvalidId:
        raise HTTPException(status_code=400, detail="Bad id")
    # ownership is enforced inside the query, so another user's id simply 404s
    d = await db.documents.find_one({"_id": oid, "user_id": user["_id"]})
    if not d:
        raise HTTPException(status_code=404, detail="Document not found")
    return d


@app.get("/api/documents/{doc_id}")
async def get_document(doc_id: str, user=Depends(current_user)):
    d = await _owned_doc(doc_id, user)
    return {
        "id": str(d["_id"]),
        "title": d.get("title"),
        "filename": d.get("filename"),
        "content_type": d.get("content_type"),
        "size": d.get("size", 0),
        "uploaded_at": d["uploaded_at"].isoformat() if d.get("uploaded_at") else None,
        "has_extracted_text": bool(d.get("extracted_text")),
        "extracted_text": d.get("extracted_text"),
        "extraction_status": d.get("extraction_status", "pending"),
        "extraction_error": d.get("extraction_error"),
    }


@app.post("/api/documents/{doc_id}/extract-text")
async def extract_document_text(doc_id: str, user=Depends(current_user)):
    d = await _owned_doc(doc_id, user)
    stream = await bucket.open_download_stream(d["gridfs_id"])
    raw = await stream.read()

    try:
        async with extraction_semaphore:
            extracted_text = await _extract_with_timeout(
                d.get("filename"), raw, d.get("content_type")
            )
        extraction_error = None
    except Exception as exc:
        extracted_text = ""
        extraction_error = str(exc)
        logger.exception("Text extraction retry failed for document %s", doc_id)

    ext = _get_file_ext(d.get("filename"))
    supported_format = (
        ext in {".pdf", ".docx", ".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"}
        or (d.get("content_type") or "").lower() == "application/pdf"
        or (d.get("content_type") or "").lower().startswith("image/")
    )
    extraction_status = (
        "failed" if extraction_error else
        "completed" if extracted_text else
        "no_text" if supported_format else
        "unsupported"
    )
    await db.documents.update_one(
        {"_id": d["_id"], "user_id": user["_id"]},
        {
            "$set": {
                "extracted_text": extracted_text,
                "extraction_status": extraction_status,
                "extraction_error": extraction_error,
                "extraction_completed_at": datetime.now(timezone.utc),
            }
        },
    )
    return {
        "id": str(d["_id"]),
        "extracted_text": extracted_text,
        "extraction_status": extraction_status,
        "extraction_error": extraction_error,
    }


@app.get("/api/documents/{doc_id}/download")
async def download(doc_id: str, user=Depends(current_user)):
    d = await _owned_doc(doc_id, user)
    stream = await bucket.open_download_stream(d["gridfs_id"])

    async def it():
        while True:
            chunk = await stream.readchunk()
            if not chunk:
                break
            yield chunk

    filename = (d.get("filename") or "file").replace('"', "")
    return StreamingResponse(
        it(),
        media_type=d.get("content_type") or "application/octet-stream",
        headers={"Content-Disposition": 'attachment; filename="' + filename + '"'},
    )


@app.delete("/api/documents/{doc_id}")
async def delete_document(doc_id: str, user=Depends(current_user)):
    d = await _owned_doc(doc_id, user)
    try:
        await bucket.delete(d["gridfs_id"])
    except Exception:
        pass
    await db.documents.delete_one({"_id": d["_id"], "user_id": user["_id"]})
    return {"ok": True}


@app.exception_handler(HTTPException)
async def http_exc(request: Request, exc: HTTPException):
    if exc.status_code == 401 and not request.url.path.startswith("/api/"):
        return RedirectResponse("/login")
    return JSONResponse({"detail": exc.detail}, status_code=exc.status_code)
